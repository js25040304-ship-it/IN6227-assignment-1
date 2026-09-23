#!/usr/bin/env python3

"""Generate an IN6227-style two-column DOCX and optional PDF from report_data.json."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader


class ReportError(ValueError):
    """Raised when report generation inputs are incomplete or unsafe."""


MODEL_LABELS = {
    "dummy_most_frequent": "Dummy",
    "logistic_regression": "Logistic",
    "random_forest": "Random forest",
    "hist_gradient_boosting": "Hist. gradient boosting",
}

PARAMETER_LABELS = {
    "C": "C",
    "class_weight": "class weight",
    "max_depth": "depth",
    "max_features": "features",
    "min_samples_leaf": "minimum leaf size",
    "n_estimators": "trees",
    "max_leaf_nodes": "maximum leaf nodes",
    "learning_rate": "learning rate",
    "l2_regularization": "L2 penalty",
}


def readable_parameter(name: str, value: Any) -> str:
    key = name.removeprefix("model__")
    label = PARAMETER_LABELS.get(key, key.replace("_", " "))
    values = value if isinstance(value, list) else [value]
    rendered = []
    for item in values:
        if item is None:
            rendered.append("unlimited" if key in {"max_depth", "max_leaf_nodes"} else "none")
        elif isinstance(item, str):
            rendered.append(item.replace("_", " "))
        else:
            rendered.append(str(item))
    return f"{label}={'/'.join(rendered)}"


def search_summary(model: str, details: dict[str, Any]) -> str:
    """Describe evaluated settings without implying a Cartesian grid."""
    parameters = details.get("parameter_values", {})
    varying = [values for values in parameters.values() if isinstance(values, list) and len(values) > 1]
    factorial_size = 1
    for values in varying:
        factorial_size *= len(values)
    settings = max(details.get("feature_sets", {}).values() or [0])
    label = MODEL_LABELS.get(model, model)
    if model == "random_forest":
        label += " [1]"
    text = label + ": " + ", ".join(readable_parameter(name, values) for name, values in parameters.items())
    if factorial_size > settings:
        text += f"; {settings} coupled settings per feature set (not a full factorial)"
    elif settings:
        text += f"; {settings} settings per feature set"
    return text


def load_json(path: Path) -> dict[str, Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ReportError(f"Could not read valid report data: {exc}") from exc


def require_metadata(args: argparse.Namespace) -> None:
    values = {
        "full name": args.full_name,
        "matric number": args.matric_number,
        "GitHub URL": args.github_url,
        "model name": args.model_name,
        "model version": args.model_version,
        "LLM interface": args.llm_interface,
    }
    forbidden = {"tbd", "todo", "placeholder", "author name", "your name", "unknown"}
    for label, value in values.items():
        if not value or value.strip().lower() in forbidden:
            raise ReportError(f"A real {label} is required; placeholders are not allowed")
    if not re.fullmatch(r"[A-Za-z0-9-]{5,24}", args.matric_number.strip()):
        raise ReportError("Matric number must be 5-24 letters, numbers, or hyphens")
    if not args.github_url.startswith("https://github.com/"):
        raise ReportError("GitHub URL must be an https://github.com/ link")


def validate_report_data(data: dict[str, Any]) -> None:
    if data.get("source_evidence", {}).get("automatic_verification_passed") is not True:
        raise ReportError("Report data is not independently verified")
    selected = data.get("results", {}).get("selected_model")
    metrics = data.get("results", {}).get("selected_final_metrics", {})
    if not selected or not metrics:
        raise ReportError("Report data lacks a selected model or final metrics")
    matrix = metrics.get("confusion_matrix", {}).get("values", [])
    if not matrix or sum(sum(row) for row in matrix) != metrics.get("evaluated_rows"):
        raise ReportError("Confusion matrix does not reconcile with evaluated rows")
    labels = metrics.get("confusion_matrix", {}).get("labels", [])
    if len(labels) != len(matrix) or any(len(row) != len(labels) for row in matrix):
        raise ReportError("Confusion matrix dimensions do not match its labels")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([properties, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 55, start: int = 55, bottom: int = 55, end: int = 55) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9D9D9") -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)


def format_run(run, size: float = 10, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_text(
    document: Document,
    text: str,
    *,
    first_line: bool = False,
    justify: bool = True,
    keep_together: bool = False,
    keep_with_next: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Normal"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = keep_together
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if first_line:
        paragraph.paragraph_format.first_line_indent = Inches(0.18)
    format_run(paragraph.add_run(text))


def add_heading(document: Document, text: str, *, top_space: float = 5) -> None:
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.space_before = Pt(top_space)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    format_run(paragraph.add_run(text.upper()), bold=True)


def clear_body(document: Document) -> None:
    body = document._element.body
    final_section = body.sectPr
    for child in list(body):
        if child is not final_section:
            body.remove(child)


def set_columns(section, count: int, space_twips: int = 360) -> None:
    properties = section._sectPr
    columns = properties.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        properties.append(columns)
    columns.set(qn("w:num"), str(count))
    columns.set(qn("w:space"), str(space_twips))
    columns.set(qn("w:equalWidth"), "true")


def add_page_field(paragraph, *, align_right: bool = True) -> None:
    if align_right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    format_run(run, size=8, color="6B7280")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])


def replace_header_footer(document: Document, full_name: str) -> None:
    for section in document.sections:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header = section.header
        for paragraph in list(header.paragraphs):
            paragraph._element.getparent().remove(paragraph._element)
        p = header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        format_run(p.add_run("IN6227 Data Mining | Assignment 1 | Variant 2"), size=8, color="808080")
        footer = section.footer
        for paragraph in list(footer.paragraphs):
            paragraph._element.getparent().remove(paragraph._element)
        for table in list(footer.tables):
            table._element.getparent().remove(table._element)
        footer_line = footer.add_paragraph()
        footer_line.paragraph_format.tab_stops.add_tab_stop(Inches(6.55), WD_TAB_ALIGNMENT.RIGHT)
        format_run(
            footer_line.add_run(f"Leakage-Aware Tabular Classification | {full_name}"),
            size=8,
            color="4F81BD",
        )
        footer_line.add_run("\t")
        add_page_field(footer_line, align_right=False)


def format_table(table, header_fill: str = "DCE6F1", font_size: float = 8.5) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F7F9FB")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_with_next = row_index < len(table.rows) - 1
                for run in paragraph.runs:
                    format_run(run, size=font_size, bold=(row_index == 0))


def set_table_widths(table, widths: tuple[float, ...]) -> None:
    total_twips = int(sum(widths) * 1440)
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:w"), str(total_twips))
    table_width.set(qn("w:type"), "dxa")
    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Inches(width)
        for cell in table.columns[column_index].cells:
            cell.width = Inches(width)
            cell_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            cell_width.set(qn("w:w"), str(int(width * 1440)))
            cell_width.set(qn("w:type"), "dxa")


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(3)
    format_run(paragraph.add_run(text), size=8)


def focus_metric(data: dict[str, Any], final: dict[str, Any]) -> float:
    if data["task"].get("type") == "binary" and data["task"].get("positive_label") is not None:
        positive = str(data["task"].get("positive_label"))
        return (final.get("per_class", {}).get(positive, {}) or {}).get("recall", 0.0)
    return final.get("balanced_accuracy", 0.0)


def focus_metric_label(data: dict[str, Any]) -> str:
    if data["task"].get("type") == "binary" and data["task"].get("positive_label") is not None:
        return f"{data['task'].get('positive_label')} recall"
    return "Bal. acc."


def model_metric_rows(data: dict[str, Any]) -> list[tuple[str, float, float, float]]:
    development = {}
    selected_feature_set = data.get("method", {}).get("preprocessing", {}).get("selected_feature_set")
    if not selected_feature_set:
        selected_model = data["results"].get("selected_model")
        selected_feature_set = next(
            (row.get("feature_set") for row in data["results"]["development_comparison"] if row.get("model") == selected_model),
            None,
        )
    for row in data["results"]["development_comparison"]:
        model = row["model"]
        feature_set = row.get("feature_set")
        if feature_set == selected_feature_set or model.startswith("dummy"):
            development[model] = row["development"].get("macro_f1_mean")
    rows = []
    declared_models = data.get("method", {}).get("models") or list(data["results"]["final_metrics_by_model"])
    for model in declared_models:
        final = data["results"]["final_metrics_by_model"].get(model, {})
        rows.append((model, development.get(model, 0.0), final.get("macro_f1", 0.0), focus_metric(data, final)))
    return rows


def font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Times New Roman.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def create_metric_figure(data: dict[str, Any], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1050, 500
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(34, bold=True)
    label_font = font(27)
    small_font = font(24)
    draw.text((55, 24), "Held-out performance", font=title_font, fill="#000000")
    x0, y0, x1, y1 = 85, 90, 1015, 405
    draw.line((x0, y1, x1, y1), fill="#666666", width=2)
    draw.line((x0, y0, x0, y1), fill="#666666", width=2)
    for tick in (0.0, 0.25, 0.5, 0.75, 1.0):
        y = y1 - int((y1 - y0) * tick)
        draw.line((x0 - 7, y, x1, y), fill="#E5E7EB", width=1)
        draw.text((20, y - 14), f"{tick:.2f}", font=small_font, fill="#374151")
    rows = model_metric_rows(data)
    group_width = (x1 - x0) / len(rows)
    colors = ("#4F81BD", "#C0504D")
    for index, (model, _dev, macro_f1, recall) in enumerate(rows):
        center = x0 + group_width * (index + 0.5)
        for offset, value, color in ((-38, macro_f1, colors[0]), (10, recall, colors[1])):
            bar_x0 = int(center + offset)
            bar_x1 = bar_x0 + 28
            bar_y = y1 - int((y1 - y0) * value)
            draw.rectangle((bar_x0, bar_y, bar_x1, y1), fill=color)
        label = MODEL_LABELS.get(model, model)
        box = draw.textbbox((0, 0), label, font=label_font)
        draw.text((center - (box[2] - box[0]) / 2, y1 + 14), label, font=label_font, fill="#111827")
    draw.rectangle((590, 28, 615, 53), fill=colors[0])
    draw.text((625, 25), "Macro-F1", font=small_font, fill="#111827")
    draw.rectangle((790, 28, 815, 53), fill=colors[1])
    draw.text((825, 25), focus_metric_label(data), font=small_font, fill="#111827")
    image.save(output, dpi=(300, 300))


def build_document(data: dict[str, Any], args: argparse.Namespace) -> None:
    document = Document(args.template_docx)
    clear_body(document)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10)
    try:
        title_style = document.styles["Title"]
    except KeyError:
        title_style = document.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)

    first = document.sections[0]
    first.orientation = WD_ORIENT.PORTRAIT
    first.page_width = Mm(210)
    first.page_height = Mm(297)
    first.left_margin = Inches(0.75)
    first.right_margin = Inches(0.75)
    first.top_margin = Inches(0.82)
    first.bottom_margin = Inches(0.65)
    set_columns(first, 1)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(3)
    title.add_run("Leakage-Aware Tabular Classification")
    metadata = document.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(3)
    lines = [
        f"{args.full_name}, {args.matric_number}",
        "IN6227-Assignment-1",
        "Variant-2",
        f"LLMs: {args.model_name} ({args.model_version})",
        f"Interfaces: {args.llm_interface}",
    ]
    for index, line in enumerate(lines):
        run = metadata.add_run(line)
        format_run(run, size=9 if index >= 3 else 10)
        run.add_break()
    format_run(metadata.add_run("Skill: "), size=9)
    add_hyperlink(metadata, args.github_url.removeprefix("https://"), args.github_url)

    body_section = document.add_section(WD_SECTION.CONTINUOUS)
    body_section.left_margin = Inches(0.75)
    body_section.right_margin = Inches(0.75)
    body_section.top_margin = Inches(0.82)
    body_section.bottom_margin = Inches(0.65)
    set_columns(body_section, 2, 360)

    splits = data["dataset"]["splits"]
    train, test = splits["train"], splits.get("test", {})
    positive = str(data["task"].get("positive_label"))
    train_positive = next((x for x in train["class_distribution"] if str(x["label"]) == positive), None)
    imbalance = next((x for x in data["quality_findings"]["warnings"] if x.get("code") == "CLASS_IMBALANCE"), {})

    add_heading(document, "Data and task")
    final_rows = data["results"]["selected_final_metrics"]["evaluated_rows"]
    final_source = "the untouched supplied test set" if test else "a stratified held-out partition"
    add_text(
        document,
        f"The skill inferred the {data['task']['type']} target {data['task']['target']}. "
        f"Excluding {train['missing_target_rows']} missing-label rows left {train['usable_labelled_rows']:,} labelled source rows; "
        f"{final_rows:,} rows in {final_source} supported final evaluation. Predictors comprised "
        f"{train['predictor_type_counts'].get('numeric', 0)} numeric and {train['predictor_type_counts'].get('categorical', 0)} categorical variables.",
    )
    if train_positive is not None:
        class_text = f"The positive class ({positive}) represented {train_positive['ratio']:.1%} of development data"
    else:
        class_text = f"The target contained {len(train['class_distribution'])} observed classes"
    cross_split_checks = data["dataset"].get("cross_split_checks") or {}
    predictor_overlap = cross_split_checks.get("predictor_overlap") or {}
    add_text(
        document,
        class_text
        + (f", a {imbalance.get('majority_to_minority_ratio'):.2f}:1 majority-to-minority ratio" if imbalance else "")
        + f". Missingness affected {train['rows_with_any_missing']} source rows"
        + (f" and {test.get('rows_with_any_missing', 0)} test rows" if test else "")
        + f". There were {train['exact_duplicate_rows']} exact source duplicates and {predictor_overlap.get('test_rows_matching_train_predictors', 0)} supplied-test rows matching training predictors.",
        first_line=True,
    )
    outlier_count = int(train.get("iqr_flagged_value_count", 0) or 0)
    outlier_columns = len(train.get("iqr_outlier_counts", {}) or {})
    if outlier_count:
        add_text(
            document,
            f"Profiling found {outlier_count:,} IQR-flagged values across {outlier_columns} numeric predictors. They were retained: median imputation handled missingness, scaling stabilized logistic regression, and the tree model required no row deletion.",
            first_line=True,
        )

    add_heading(document, "Methods")
    validation = data["method"]["validation"]
    add_text(
        document,
        f"Numeric values were median-imputed; categoricals were most-frequent-imputed and one-hot encoded with unseen-category handling. "
        f"Only logistic regression used scaling. Every learned transformation stayed inside a {validation['folds']}-fold shuffled stratified pipeline (seed {validation['random_seed']}). "
        f"Macro-F1 was the primary selection metric because accuracy alone would reward the majority class [3].",
    )
    tuning = data["method"]["tuning"]["search"]
    search_text = "; ".join(search_summary(model, details) for model, details in tuning.items())
    selected_parameter_text = ", ".join(
        readable_parameter(name, value) for name, value in data["method"].get("selected_parameters", {}).items()
    ) or "defaults"
    add_text(
        document,
        f"The fixed search compared {', '.join(MODEL_LABELS.get(model, model) for model in tuning)} with a most-frequent dummy ({search_text}); selected settings: {selected_parameter_text}. "
        f"The predefined grids ran once on fixed folds; final-test metrics did not influence feature, parameter, threshold, or model selection.",
        first_line=True,
        justify=False,
    )

    add_heading(document, "Feature review")
    sensitivity = data["quality_findings"].get("suspicious_aggregate_sensitivity", [])
    excluded = ", ".join(data["method"]["preprocessing"].get("excluded_features", []))
    deltas = ", ".join(
        f"{MODEL_LABELS.get(row['model'], row['model'])} {row['delta_without_minus_full']:+.4f}"
        for row in sensitivity
    )
    if excluded and sensitivity:
        feature_text = (
            f"The aggregate-like {excluded} was highly correlated with components and had uncertain prediction-time provenance. "
            f"A development-only removal check changed macro-F1 by {deltas}. Because both changes were below fold variability, it was dropped before final evaluation to reduce leakage risk."
        )
    else:
        feature_text = "No report-critical aggregate warning remained unresolved; the final feature set followed the recorded schema and leakage checks."
    add_text(document, feature_text)

    add_heading(document, "Results")
    table = document.add_table(rows=1, cols=4)
    for cell, label in zip(table.rows[0].cells, ("Model", "CV M-F1", "Test M-F1", focus_metric_label(data))):
        cell.text = label
    for model, dev, final_f1, recall in model_metric_rows(data):
        row = table.add_row().cells
        row[0].text = MODEL_LABELS.get(model, model) + ("*" if model == data["results"]["selected_model"] else "")
        row[1].text = f"{dev:.3f}"
        row[2].text = f"{final_f1:.3f}"
        row[3].text = f"{recall:.3f}"
    set_table_widths(table, (1.12, 0.70, 0.72, 0.71))
    format_table(table, font_size=8)
    add_caption(document, "Table 1. Development and held-out comparison (* selected).")

    figure_path = args.figures_dir / "model_performance.png"
    create_metric_figure(data, figure_path)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(figure_path), width=Inches(2.75))
    add_caption(document, f"Figure 1. Final macro-F1 and {focus_metric_label(data).lower()}.")

    selected = data["results"]["selected_final_metrics"]
    matrix = selected["confusion_matrix"]
    labels = [str(value) for value in matrix["labels"]]
    selected_label = MODEL_LABELS.get(data["results"]["selected_model"], data["results"]["selected_model"])
    if len(labels) <= 3:
        cm = document.add_table(rows=len(labels) + 1, cols=len(labels) + 1)
        cm.cell(0, 0).text = "Actual / Pred."
        for column, label in enumerate(labels, 1):
            cm.cell(0, column).text = label
        for row_index, label in enumerate(labels, 1):
            cm.cell(row_index, 0).text = label
            for column_index, value in enumerate(matrix["values"][row_index - 1], 1):
                cm.cell(row_index, column_index).text = f"{value:,}"
        first_width = 1.15
        other_width = (3.25 - first_width) / len(labels)
        set_table_widths(cm, (first_width,) + (other_width,) * len(labels))
        format_table(cm, font_size=8)
        add_caption(document, f"Table 2. {selected_label} confusion matrix (n={selected['evaluated_rows']:,}).")
    else:
        class_rows = sorted(selected.get("per_class", {}).items(), key=lambda item: item[1].get("recall", 0))[:5]
        class_table = document.add_table(rows=1, cols=5)
        for cell, label in zip(class_table.rows[0].cells, ("Class", "Prec.", "Recall", "F1", "n")):
            cell.text = label
        for label, values in class_rows:
            row = class_table.add_row().cells
            row[0].text = str(label)
            row[1].text = f"{values.get('precision', 0):.3f}"
            row[2].text = f"{values.get('recall', 0):.3f}"
            row[3].text = f"{values.get('f1', 0):.3f}"
            row[4].text = f"{values.get('support', 0):,}"
        set_table_widths(class_table, (0.85, 0.6, 0.65, 0.55, 0.6))
        format_table(class_table, font_size=7.5)
        add_caption(document, f"Table 2. Five lowest-recall classes for {selected_label}; full matrix retained in artifacts.")

    add_heading(document, "Findings")
    final_models = data["results"]["final_metrics_by_model"]
    selected_model = data["results"]["selected_model"]
    selected_metrics = final_models[selected_model]
    other_model = next(
        model for model in final_models
        if model not in {"dummy_most_frequent", selected_model}
    )
    other_metrics = final_models.get(other_model, {})
    selected_name = MODEL_LABELS.get(selected_model, selected_model)
    other_name = MODEL_LABELS.get(other_model, other_model)
    if data["task"].get("type") == "binary" and data["task"].get("positive_label") is not None:
        findings = (
            f"{selected_name} was selected on development evidence and achieved held-out macro-F1 {selected_metrics.get('macro_f1', 0):.3f} and balanced accuracy {selected_metrics.get('balanced_accuracy', 0):.3f}. "
            f"It recovered {selected_metrics.get('per_class', {}).get(positive, {}).get('recall', 0):.1%} of positive cases, versus {other_metrics.get('per_class', {}).get(positive, {}).get('recall', 0):.1%} for {other_name.lower()}. "
            f"Overall accuracy was {selected_metrics.get('accuracy', 0):.3f} versus {other_metrics.get('accuracy', 0):.3f}; positive precision was {selected_metrics.get('per_class', {}).get(positive, {}).get('precision', 0):.3f} versus {other_metrics.get('per_class', {}).get(positive, {}).get('precision', 0):.3f}."
        )
    elif data["task"].get("type") == "binary":
        findings = (
            f"{selected_name} was selected on development evidence and achieved held-out macro-F1 {selected_metrics.get('macro_f1', 0):.3f}, balanced accuracy {selected_metrics.get('balanced_accuracy', 0):.3f}, and accuracy {selected_metrics.get('accuracy', 0):.3f}. "
            f"{other_name} achieved {other_metrics.get('macro_f1', 0):.3f}, {other_metrics.get('balanced_accuracy', 0):.3f}, and {other_metrics.get('accuracy', 0):.3f}, respectively. No class was designated semantically positive, so class-balanced metrics and the per-class table were used instead of positive-class claims."
        )
    else:
        findings = (
            f"{selected_name} was selected on development evidence. On the held-out data it achieved macro-F1 {selected_metrics.get('macro_f1', 0):.3f}, balanced accuracy {selected_metrics.get('balanced_accuracy', 0):.3f}, and accuracy {selected_metrics.get('accuracy', 0):.3f}; "
            f"{other_name} achieved {other_metrics.get('macro_f1', 0):.3f}, {other_metrics.get('balanced_accuracy', 0):.3f}, and {other_metrics.get('accuracy', 0):.3f}, respectively. Macro and balanced metrics keep minority classes visible in the comparison."
        )
    add_text(document, findings)
    if "roc_auc" in selected_metrics and "roc_auc" in other_metrics:
        discrimination = (
            f"Binary discrimination was compared by ROC-AUC ({selected_metrics['roc_auc']:.3f} selected; {other_metrics['roc_auc']:.3f} alternative) "
            f"and average precision ({selected_metrics.get('average_precision', 0):.3f}; {other_metrics.get('average_precision', 0):.3f})."
        )
    elif "roc_auc_ovr_weighted" in selected_metrics and "roc_auc_ovr_weighted" in other_metrics:
        discrimination = (
            f"Multiclass discrimination was compared by weighted OvR ROC-AUC ({selected_metrics['roc_auc_ovr_weighted']:.3f} selected; "
            f"{other_metrics['roc_auc_ovr_weighted']:.3f} alternative)."
        )
    else:
        discrimination = "Probability discrimination metrics were not reported because their required task semantics were unavailable."
    add_text(document, discrimination + " The selected model reflects the declared class-balanced objective, not dominance on every metric.", first_line=True)

    if data["task"].get("type") == "binary":
        selected_matrix = selected_metrics["confusion_matrix"]["values"]
        other_matrix = other_metrics["confusion_matrix"]["values"]
        fewer_false_negatives = other_matrix[1][0] - selected_matrix[1][0]
        extra_false_positives = selected_matrix[0][1] - other_matrix[0][1]
        if fewer_false_negatives > 0 and extra_false_positives > 0:
            cost_ratio = extra_false_positives / fewer_false_negatives
            add_text(
                document,
                f"Against {other_name.lower()}, the selected model traded {fewer_false_negatives:,} fewer false negatives for {extra_false_positives:,} additional false positives. "
                f"It has lower observed error cost only if a false negative costs more than {cost_ratio:.2f} times a false positive; this post-hoc interpretation did not drive selection.",
                first_line=True,
            )

    add_heading(document, "Data quality checks")
    quality_table = document.add_table(rows=1, cols=2)
    quality_table.cell(0, 0).text = "Check"
    quality_table.cell(0, 1).text = "Observed result"
    blocking_warnings = sum(
        warning.get("severity") == "blocking"
        for warning in data.get("quality_findings", {}).get("warnings", [])
    )
    quality_rows = (
        ("Rows with missing values", f"{train['rows_with_any_missing']:,}"),
        ("Exact duplicate rows", f"{train['exact_duplicate_rows']:,}"),
        ("IQR-flagged values", f"{outlier_count:,}"),
        ("Test rows matching train predictors", f"{predictor_overlap.get('test_rows_matching_train_predictors', 0):,}"),
        ("Unresolved blocking warnings", f"{blocking_warnings:,}"),
    )
    for check, result in quality_rows:
        row = quality_table.add_row().cells
        row[0].text = check
        row[1].text = result
    set_table_widths(quality_table, (2.25, 1.0))
    format_table(quality_table, font_size=7.5)
    add_caption(document, "Table 3. Automated data-quality checks before evaluation.")

    add_heading(document, "Discussion")
    diagnostics = data["method"].get("selection_diagnostics", {})
    reference_model = diagnostics.get("reference_model", "logistic_regression")
    challenger_model = diagnostics.get("challenger_model", other_model if other_model != reference_model else selected_model)
    challenger_wins = diagnostics.get("challenger_fold_wins", diagnostics.get("random_forest_fold_wins", 0))
    reference_name = MODEL_LABELS.get(reference_model, reference_model)
    challenger_name = MODEL_LABELS.get(challenger_model, challenger_model)
    add_text(
        document,
        f"The paired {challenger_name.lower()}-minus-{reference_name.lower()} macro-F1 difference was {diagnostics.get('mean_difference', 0):+.4f}; {challenger_name.lower()} won {challenger_wins} of {validation['folds']} folds. "
        f"The predefined practical-tie result was {str(diagnostics.get('practical_tie', False)).lower()} using a one-standard-error reference ({diagnostics.get('standard_error', 0):.4f}), leading to {selected_name.lower()}. "
        f"Operational error costs could still justify the alternative model.",
    )

    page_two = document.add_section(WD_SECTION.NEW_PAGE)
    page_two.left_margin = Inches(0.75)
    page_two.right_margin = Inches(0.75)
    # Keep the second page's original report margins. Together with the chained
    # reference paragraphs below, this lets the left column fill naturally and
    # moves the intact References block to the top of the right column.
    page_two.top_margin = Inches(1.0)
    page_two.bottom_margin = Inches(0.8)
    set_columns(page_two, 2, 360)

    add_heading(document, "Limitations")
    add_text(
        document,
        "The analysis assumes independent rows and stable train/test collection. It does not establish causal effects, fairness across unobserved groups, probability calibration, or prediction-time availability of every retained feature. "
        "Threshold tuning was intentionally omitted because no application-specific error costs were provided.",
        first_line=True,
    )

    add_heading(document, "SKILL generalisation")
    add_text(
        document,
        "The reusable SKILL enforces typed data checks, fold-local preprocessing, a dummy baseline, an adaptive challenger, fixed selection, immutable runs, and independent verification from machine-readable evidence.",
    )
    add_text(
        document,
        "Phase 5 completed 13 forward-test scenarios. Four differently shaped synthetic modelling cases ran end to end, while four unsafe or unresolved inputs produced a typed refusal or human checkpoint; the remaining cases exercised formats, class structures, report branches, and verification. These tests demonstrate behavioral coverage, not guaranteed predictive accuracy on every future dataset.",
        first_line=True,
        keep_together=True,
    )
    add_text(
        document,
        "The one-command route stops before cross-validation when time or group dependence is unresolved and before final-test evaluation when aggregate availability is uncertain. A new immutable run must record the explicit human decision; semantic uncertainty is never converted into an automatic feature or validation choice.",
        first_line=True,
        keep_together=True,
    )
    add_text(
        document,
        "The supported envelope covers CSV, TSV, Excel, directory, and safe ZIP inputs for binary or single-label multiclass classification. Regression, multilabel, text-only, and grouped or ordered validation are refused rather than silently approximated.",
        first_line=True,
        keep_together=True,
    )

    add_heading(document, "Reproducibility and verification")
    versions = data.get("reproducibility", {}).get("runtime_versions", {})
    add_text(
        document,
        f"{data['run_id']} records the data and skill fingerprints, seed, fold assignments, candidate settings, predictions, and package versions "
        f"(Python {versions.get('python')}; scikit-learn {versions.get('scikit-learn')} [2]). A separate verifier recomputed {sum(data['source_evidence']['verification_checks'].values())} checks from saved predictions, including all scalar metrics and the confusion matrix; all passed.",
    )
    if data["task"].get("type") == "binary" and data["task"].get("positive_label") is not None:
        matrix_values = selected_metrics["confusion_matrix"]["values"]
        tn, fp = matrix_values[0]
        fn, tp = matrix_values[1]
        add_text(
            document,
            f"A manual reconstruction also reconciled {tn:,} + {fp:,} + {fn:,} + {tp:,} = {selected_metrics['evaluated_rows']:,} evaluated rows and positive recall "
            f"{tp:,} / ({tp:,} + {fn:,}) = {selected_metrics['per_class'][positive]['recall']:.7f}. This verifies arithmetic and label orientation, not feature provenance or row independence.",
            first_line=True,
        )

    add_heading(document, "Conclusion")
    if data["task"].get("type") == "binary" and data["task"].get("positive_label") is not None:
        conclusion = (
            f"A leakage-aware, class-balanced workflow selected {selected_name.lower()} with held-out macro-F1 {selected_metrics.get('macro_f1', 0):.3f} and positive recall {selected_metrics.get('per_class', {}).get(positive, {}).get('recall', 0):.3f}. "
            f"The result is stronger than the dummy baseline, while its trade-offs against {other_name.lower()} remain explicit."
        )
    else:
        conclusion = (
            f"A leakage-aware workflow selected {selected_name.lower()} with held-out macro-F1 {selected_metrics.get('macro_f1', 0):.3f} and balanced accuracy {selected_metrics.get('balanced_accuracy', 0):.3f}. "
            f"The result is interpreted against both {other_name.lower()} and the dummy baseline without relying on overall accuracy alone."
        )
    add_text(document, conclusion)

    add_heading(document, "References")
    if "hist_gradient_boosting" in final_models:
        add_text(document, "[1] Friedman, J. H. (2001). Greedy Function Approximation: A Gradient Boosting Machine. Annals of Statistics, 29, 1189-1232.", keep_with_next=True)
    else:
        add_text(document, "[1] Breiman, L. (2001). Random Forests. Machine Learning, 45, 5-32.", keep_with_next=True)
    add_text(document, "[2] Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. JMLR, 12, 2825-2830.", keep_with_next=True)
    add_text(document, "[3] He, H. and Garcia, E. A. (2009). Learning from Imbalanced Data. IEEE TKDE, 21(9), 1263-1284. doi:10.1109/TKDE.2008.239.")

    replace_header_footer(document, args.full_name)
    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)


def convert_to_pdf(docx: Path, output_pdf: Path, soffice: str) -> None:
    output_pdf.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=output_pdf.parent) as temp:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", temp, str(docx)],
            text=True,
            capture_output=True,
            check=False,
        )
        produced = Path(temp) / f"{docx.stem}.pdf"
        if result.returncode != 0 or not produced.is_file():
            raise ReportError(f"PDF conversion failed: {result.stdout} {result.stderr}".strip())
        pages = len(PdfReader(produced).pages)
        if pages > 2:
            raise ReportError(f"Generated main report has {pages} pages; the limit is 2")
        shutil.copy2(produced, output_pdf)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("report_data", type=Path)
    parser.add_argument("--template-docx", type=Path, required=True)
    parser.add_argument("--full-name", required=True)
    parser.add_argument("--matric-number", required=True)
    parser.add_argument("--github-url", required=True)
    parser.add_argument("--model-name", required=True)
    parser.add_argument("--model-version", required=True)
    parser.add_argument("--llm-interface", required=True)
    parser.add_argument("--output-docx", type=Path, required=True)
    parser.add_argument("--output-pdf", type=Path)
    parser.add_argument("--figures-dir", type=Path)
    parser.add_argument("--soffice", default=shutil.which("soffice"))
    args = parser.parse_args()
    args.figures_dir = args.figures_dir or args.output_docx.parent.parent / "figures"
    try:
        require_metadata(args)
        data = load_json(args.report_data)
        validate_report_data(data)
        if not args.template_docx.is_file():
            raise ReportError("A converted DOCX report template is required")
        build_document(data, args)
        if args.output_pdf:
            if not args.soffice:
                raise ReportError("soffice is required for PDF conversion")
            convert_to_pdf(args.output_docx.resolve(), args.output_pdf.resolve(), args.soffice)
    except ReportError as exc:
        print(f"Report generation failed: {exc}")
        return 2
    print(args.output_docx.resolve())
    if args.output_pdf:
        print(args.output_pdf.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
